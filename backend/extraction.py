from dataclasses import dataclass
from pathlib import Path

import pytesseract
from docx import Document as DocxDocument
from pdf2image import convert_from_path
from PIL import Image
from pypdf import PdfReader

# Pages with less text than this are treated as scanned images
MIN_TEXT_LENGTH = 20
OCR_DPI = 300
OCR_LANGUAGE = "eng"


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    method: str  # "text" or "ocr"


def ocr_image(image: Image.Image) -> str:
    return pytesseract.image_to_string(image, lang=OCR_LANGUAGE).strip()


def extract_pdf(file_path: Path) -> list[ExtractedPage]:
    reader = PdfReader(file_path)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if len(text) >= MIN_TEXT_LENGTH:
            pages.append(ExtractedPage(page_number, text, "text"))
            continue

        # No real text on this page, so it's probably a scan: render it and OCR it
        images = convert_from_path(
            file_path,
            dpi=OCR_DPI,
            first_page=page_number,
            last_page=page_number,
        )
        pages.append(ExtractedPage(page_number, ocr_image(images[0]), "ocr"))

    return pages


def extract_docx(file_path: Path) -> list[ExtractedPage]:
    docx = DocxDocument(file_path)
    lines = [paragraph.text for paragraph in docx.paragraphs if paragraph.text.strip()]

    # Text inside tables isn't part of docx.paragraphs, so read it separately
    for table in docx.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    # DOCX files have no fixed pages, so everything is stored as page 1
    return [ExtractedPage(1, "\n".join(lines), "text")]


def extract_image(file_path: Path) -> list[ExtractedPage]:
    with Image.open(file_path) as image:
        return [ExtractedPage(1, ocr_image(image), "ocr")]


def extract_pages(file_path: Path, content_type: str) -> list[ExtractedPage]:
    if content_type == "application/pdf":
        return extract_pdf(file_path)

    if content_type.startswith("image/"):
        return extract_image(file_path)

    return extract_docx(file_path)
